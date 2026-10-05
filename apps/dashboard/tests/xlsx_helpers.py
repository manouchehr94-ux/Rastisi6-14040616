"""Shared helpers for the XLSX import/export tests: build upload workbooks with
plain openpyxl (independent of the application's own writer) and read generated
workbooks back through a real XLSX reader."""

import io
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_xlsx(headers, rows, *, sheet="داده‌ها", extra_sheets=None, cell_types=None) -> bytes:
    """Return the bytes of a workbook whose first sheet holds ``headers`` + ``rows``.

    ``cell_types`` maps ``(row_index, col_index)`` (1-based, header row = 1) to an
    openpyxl ``data_type`` override (e.g. ``"f"`` to force a formula cell)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    for (row_index, col_index), data_type in (cell_types or {}).items():
        ws.cell(row=row_index, column=col_index).data_type = data_type
    for title, sheet_rows in (extra_sheets or {}).items():
        extra = wb.create_sheet(title)
        for sheet_row in sheet_rows:
            extra.append(list(sheet_row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def xlsx_upload(headers, rows, *, name="data.xlsx", **kwargs) -> SimpleUploadedFile:
    return SimpleUploadedFile(name, make_xlsx(headers, rows, **kwargs), content_type=XLSX_MIME)


def open_workbook(content: bytes):
    return load_workbook(io.BytesIO(content))


def sheet_values(ws):
    """All rows of ``ws`` as lists of plain cell values."""
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def read_job_file(field_file) -> bytes:
    field_file.open("rb")
    try:
        return field_file.read()
    finally:
        field_file.close()


def zip_names(content: bytes):
    return set(zipfile.ZipFile(io.BytesIO(content)).namelist())
