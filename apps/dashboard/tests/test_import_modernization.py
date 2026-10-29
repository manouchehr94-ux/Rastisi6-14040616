"""Excel import modernization: explicit final-import step with a confirmation modal showing the real numbers, product NAME as the main
column (SKU secondary, tenant-scoped fallback), clean error wording (no «برندبرند…» duplication) and the unchanged safety rules."""

from decimal import Decimal

from django.urls import reverse

from apps.catalog.models import Brand, Product
from apps.core.models import ImportJob, ImportRowResult
from apps.dashboard.services import import_service, import_xlsx
from apps.dashboard.tests.test_import_proposals import ProposalFlowBase, row
from apps.stores.models import Store


class FinalImportModalTests(ProposalFlowBase):
    NAMES = ["آلفا", "بتا", "گاما"]

    def _file(self):
        return [row(f"کالای {n}", f"FM-{i}", brand=n) for i, n in enumerate(self.NAMES, start=1)]

    def test_button_is_disabled_until_every_proposal_is_decided_then_opens_a_modal_with_real_counts(self):
        job = self.upload(self._file())
        page = self.detail(job)
        self.assertContains(page, 'data-testid="final-import-panel"')
        self.assertContains(page, "ورود نهایی کالاها به فروشگاه")
        self.assertNotContains(page, 'id="final-import-modal"')                         # no modal (and no execute form) while blocked
        self.assertRegex(page.content.decode(), r'<button type="button" class="au-btn" disabled[^>]*data-testid="final-import-open"')
        self.assertContains(page, "را در بخش «مواردِ جدیدِ پیشنهادی» تأیید کنید")

        self.approve_all(job)
        page = self.detail(job)
        html = page.content.decode()
        self.assertIn('id="final-import-modal"', html)
        self.assertIn('data-au-open="#final-import-modal"', html)
        self.assertIn("تأیید و اجرای نهایی", html)
        self.assertIn("انصراف", html)
        self.assertIn("products.xlsx", html)
        summary = page.context["summary"]
        self.assertEqual((summary["total"], summary["will_create"], summary["will_update"], summary["invalid"], summary["awaiting"]), (3, 3, 0, 0, 0))
        counts = html[html.index('data-testid="final-import-counts"'):]
        self.assertIn("۳", counts)
        self.assertIn(reverse("dashboard:import-execute", args=[job.pk]), html)
        self.assertEqual(Product.objects.filter(store=self.store, sku__startswith="FM-").count(), 0)   # nothing happened by opening the page

    def test_double_post_of_the_final_import_executes_once(self):
        job = self.upload(self._file())
        self.approve_all(job)
        self.execute(job)
        self.execute(job)
        self.assertEqual(Product.objects.filter(store=self.store, sku__startswith="FM-").count(), 3)
        self.assertEqual(Brand.objects.filter(store=self.store, name__in=self.NAMES).count(), 3)

    def test_the_modal_is_not_offered_when_the_file_has_errors(self):
        job = self.upload([row("", "FM-9", brand="")])          # name missing → real error, not a proposal
        page = self.detail(job)
        self.assertNotContains(page, 'id="final-import-modal"')
        self.assertContains(page, "ابتدا خطاهایِ فایل را اصلاح")


class ProductNameColumnTests(ProposalFlowBase):
    def test_preview_rows_show_the_product_name_first_and_the_sku_second(self):
        job = self.upload([row("کفش ورزشی نایکی", "SHOE-1", brand="نایکی")])
        page = self.detail(job)
        self.assertContains(page, "نامِ کالا")
        html = page.content.decode()
        cell = html[html.index('data-testid="row-name"'):]
        cell = cell[:cell.index("</td>")]
        self.assertIn("کفش ورزشی نایکی", cell)
        self.assertIn("SHOE-1", cell)
        self.assertLess(cell.index("کفش ورزشی نایکی"), cell.index("SHOE-1"))

    def test_update_row_without_a_name_uses_the_existing_product_name_from_this_store(self):
        existing = Product.objects.create(store=self.store, vendor=self.vendor, category=self.leaf, name="کفش چرمی قدیمی", slug="old-shoe-mod",
                                          sku="OLD-SHOE", price=Decimal("90000"), stock=3)
        job = self.upload([["", "", "", "", existing.sku, "", "", "", ""]], mode="update_only")
        names = import_service.row_display_names(job, job.row_results.all())
        self.assertEqual(names[2]["name"], "کفش چرمی قدیمی")
        self.assertEqual(names[2]["ident"], "OLD-SHOE")

    def test_fallback_is_tenant_scoped_and_honest(self):
        other = Store.objects.create(name="دیگر", slug="import-mod-other", status=Store.Status.ACTIVE)
        from apps.catalog.models import Category, Vendor
        ov = Vendor.objects.create(store=other, name="ف", slug="ov-mod")
        oc = Category.objects.create(store=other, name="د", slug="oc-mod")
        Product.objects.create(store=other, vendor=ov, category=oc, name="کالای فروشگاهِ دیگر", slug="foreign-mod", sku="FOREIGN-1", price=Decimal("1"), stock=1)
        job = self.upload([row("", "NOPE-1")])
        fake = ImportRowResult.objects.create(import_job=job, row_number=99, source_identifier="FOREIGN-1", status="invalid", errors=["x"])
        info = import_service.row_display_names(job, [fake])[99]
        self.assertEqual(info["name"], "")                                   # another store's product is never revealed
        self.assertEqual(info["fallback"], "شناسه: FOREIGN-1")

    def test_error_report_has_a_product_name_column_and_the_column_labelled_message(self):
        job = self.upload([row("کالای خراب", "BAD-1", price="abc")])
        job.refresh_from_db()
        if job.error_report_file:
            from apps.dashboard.tests.xlsx_helpers import open_workbook, read_job_file, sheet_values
            ws = open_workbook(read_job_file(job.error_report_file))["خطاها"]
            values = sheet_values(ws)
            self.assertEqual(values[0][1], "نام کالا")
            self.assertEqual(values[1][1], "کالای خراب")
            self.assertEqual(values[1][0], 2)


class MessageWordingTests(ProposalFlowBase):
    def test_label_is_not_duplicated_next_to_the_column_chip(self):
        brand = import_xlsx.describe_errors("products", ["برندِ «نایکی» در این فروشگاه نیست و هنوز تأیید نشده است."])[0]
        self.assertEqual(brand["column"], "برند")
        self.assertTrue(brand["message"].startswith("«نایکی»"))
        sub = import_xlsx.describe_errors("products", ["زیردسته‌بندیِ «مردانه» در این فروشگاه نیست و هنوز تأیید نشده است."])[0]
        self.assertNotEqual(sub["column"], "دسته‌بندی")                      # not mapped to the MAIN category column any more
        self.assertFalse(sub["message"].startswith("زیردسته‌بندیِ"))
        main = import_xlsx.describe_errors("products", ["دسته‌بندیِ اصلیِ «پوشاک» در این فروشگاه نیست و هنوز تأیید نشده است."])[0]
        self.assertFalse(main["message"].startswith("دسته‌بندیِ"))

    def test_the_page_never_renders_the_joined_words(self):
        job = self.upload([row("پیراهن", "MW-1", category="پوشاک", sub="مردانه", brand="نایکی")])
        html = self.detail(job).content.decode()
        for glued in ("برندبرند", "دسته‌بندیزیردسته‌بندی", "دسته‌بندیدسته‌بندی"):
            self.assertNotIn(glued, html)


class QuickGuideAndVariantSampleTests(ProposalFlowBase):
    """«راهنمای سریع» on the upload page and the schema-true educational variant sample (no engine change)."""

    def test_upload_page_has_the_four_step_guide_and_every_template_link(self):
        page = self.client.get(reverse("dashboard:import-upload"))
        self.assertContains(page, 'data-testid="import-quick-guide"')
        self.assertContains(page, "راهنمای سریع ورود گروهی کالاها")
        for text in ("قالب را بگیرید", "فایل را پر و بارگذاری کنید", "مواردِ جدید را تأیید کنید", "تأیید نهایی بزنید"):
            self.assertContains(page, text)
        for kind in ("products", "variants", "inventory"):
            self.assertContains(page, reverse("dashboard:import-template", args=[kind]))
        self.assertContains(page, "با نمونه‌یِ آموزشی")

    def _variant_product(self):
        from apps.catalog.services.variant_engine_service import add_product_option

        product = Product.objects.create(store=self.store, vendor=self.vendor, category=self.leaf, name="تیشرت آموزشی", slug="tee-sample",
                                         sku="TSH-001", price=Decimal("250000"), stock=0, product_type=Product.ProductType.VARIABLE)
        add_product_option(product, label="رنگ", values=["قرمز", "آبی"])
        add_product_option(product, label="سایز", values=["L", "M"])
        return product

    def _template(self):
        from apps.dashboard.tests.xlsx_helpers import open_workbook

        response = self.client.get(reverse("dashboard:import-template", args=["variants"]))
        self.assertEqual(response.status_code, 200)
        return open_workbook(b"".join(response.streaming_content) if hasattr(response, "streaming_content") else response.content)

    def test_variant_template_has_a_sample_sheet_that_is_not_part_of_the_data_and_a_schema_true_guide(self):
        wb = self._template()
        self.assertEqual(wb.sheetnames, ["داده‌ها", "راهنما", "فهرست‌ها", "نمونه"])
        data, sample = wb["داده‌ها"], wb["نمونه"]
        self.assertEqual([r for r in data.iter_rows(min_row=2, values_only=True) if any(r)], [])  # the data sheet is empty: nothing to import by accident
        headers = [c.value for c in sample[1]]
        self.assertEqual(sample.max_row, 4)                                                   # header + 3 teaching rows
        self.assertIn("SKU کالا", headers[:2] + headers)
        guide_text = " ".join(str(c.value) for row in wb["راهنما"].iter_rows() for c in row if c.value)
        for phrase in ("ویژگی‌هایِ تنوع‌دار", "SKU تنوع", "بدونِ تغییر", "تغییر قیمت"):
            self.assertIn(phrase, guide_text)

    def test_sample_rows_run_through_the_real_variant_preview_and_execution(self):
        from io import BytesIO
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook
        from apps.catalog.models import ProductVariant

        product = self._variant_product()
        wb = load_workbook(BytesIO(self._template_bytes()))
        data, sample = wb["داده‌ها"], wb["نمونه"]
        for row in sample.iter_rows(min_row=1, values_only=True):
            data.append(list(row)) if row[0] != data.cell(1, 1).value else None
        buf = BytesIO(); wb.save(buf)
        before = ProductVariant.objects.filter(product=product).count()
        resp = self.client.post(reverse("dashboard:import-upload"), {
            "import_type": "variants", "mode": "upsert",
            "file": SimpleUploadedFile("variants-sample.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        })
        self.assertEqual(resp.status_code, 302)
        job = ImportJob.objects.filter(store=self.store).latest("pk")
        self.assertEqual(job.status, ImportJob.Status.PREVIEW_READY)
        self.assertEqual((job.total_rows, job.valid_rows, job.invalid_rows), (3, 3, 0))
        self.assertEqual(ProductVariant.objects.filter(product=product).count(), before)       # the preview changed nothing
        import_service.run_execution(job, actor=self.owner)
        job.refresh_from_db()
        self.assertEqual(job.status, ImportJob.Status.COMPLETED)
        red_l = ProductVariant.objects.get(product=product, sku="TSH-001-RED-L")
        blue_m = ProductVariant.objects.get(product=product, sku="TSH-001-BLUE-M")
        self.assertEqual((red_l.stock, blue_m.stock), (12, 8))
        self.assertEqual(blue_m.extra_price, Decimal("20000"))
        self.assertEqual(ProductVariant.objects.filter(product=product, stock=5).count(), 1)     # the third (partial) row

    def _template_bytes(self):
        response = self.client.get(reverse("dashboard:import-template", args=["variants"]))
        return b"".join(response.streaming_content) if hasattr(response, "streaming_content") else response.content
