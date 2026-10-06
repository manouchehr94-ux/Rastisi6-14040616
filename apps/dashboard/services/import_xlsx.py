"""لایه‌ی XLSXِ واردات: ستون‌هایِ فارسی، فهرست‌هایِ Store-scoped، قالبِ راهنمادار،
نگاشتِ امن به ستون‌هایِ داخلی، و گزارشِ خطا.

**اصلِ طراحی:** اعتبارسنج‌ها و سرویس‌هایِ تغییرِ داده (``import_service``)
دست‌نخورده‌اند و همچنان فقط ستون‌هایِ داخلی (``product_id``/``category_code``/
``warehouse_code``/``mode``/...) را می‌شناسند. این ماژول فقط یک *لایه‌ی نگاشت*
است: هدرِ فارسیِ فایل ← کلیدِ داخلی، و مقدارِ قابل‌خواندنِ فارسی (نامِ برند،
مسیرِ دسته‌بندی، نامِ انبار، «تنظیم موجودی نهایی»...) ← کدِ داخلی، بر پایه‌ی
رکوردهایِ *همین Store*. هیچ قاعده‌ی کسب‌وکاری اینجا تکرار نشده است.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.catalog.models import Brand, Product, ProductOption, Warehouse
from apps.catalog.services.category_path_service import normalize_path_key, store_category_paths
from apps.core.models import ImportJob
from apps.core.services import xlsx_utils
from apps.core.services.xlsx_utils import Column
from apps.core.utils import normalization_key, normalize_text
from apps.orders.models import TaxClass

MODE_LABELS = {
    ImportJob.Mode.CREATE_ONLY: "افزودن موارد جدید",
    ImportJob.Mode.UPDATE_ONLY: "فقط بروزرسانی موارد موجود",
    ImportJob.Mode.UPSERT: "افزودن و بروزرسانی",
}
MODE_DESCRIPTIONS = {
    ImportJob.Mode.CREATE_ONLY: "فقط ردیف‌هایِ تازه ساخته می‌شوند؛ اگر ردیفی با مورد موجود یکی باشد، رد می‌شود و چیزی عوض نمی‌شود.",
    ImportJob.Mode.UPDATE_ONLY: "فقط مواردِ موجود بروزرسانی می‌شوند؛ ردیفِ تازه ساخته نمی‌شود و رد می‌شود.",
    ImportJob.Mode.UPSERT: "اگر مورد وجود داشته باشد بروزرسانی می‌شود، وگرنه ساخته می‌شود. (پیشنهادی)",
}

INV_OP_SET = "تنظیم موجودی نهایی"
INV_OP_ADJUST = "افزایش/کاهش موجودی"
CLEAR_TAX_LABEL = "— بدون دسته مالیاتی —"
YES_NO = ("بله", "خیر")
SHEET_DATA = "داده‌ها"
SHEET_GUIDE = "راهنما"
SHEET_LISTS = "فهرست‌ها"
TEMPLATE_VALIDATION_ROWS = 1000


@dataclass(frozen=True)
class ImportColumn:
    key: str  # کلیدِ داخلی که اعتبارسنج می‌شناسد
    label: str  # عنوانِ ستون در فایل (فارسی)
    required: str = "no"  # always | create | no
    kind: str = "text"  # text | int | decimal | choice | bool
    help: str = ""
    accepted: str = ""
    example: str = ""
    blank: str = ""
    list_key: str | None = None  # کلیدِ فهرست در شیتِ «فهرست‌ها»
    aliases: tuple = ()
    width: float = 18
    text_format: bool = False
    in_template: bool = True

    @property
    def base_label(self) -> str:
        return _strip_label(self.label)


_PAREN_RE = re.compile(r"[\(（][^\)）]*[\)）]")


def _strip_label(label: str) -> str:
    text = _PAREN_RE.sub("", str(label or ""))
    return text.replace("*", "").strip()


def header_key(text) -> str:
    """کلیدِ مقایسه‌ی یک هدر: ستاره/پرانتز حذف، نیم‌فاصله و فاصله یکی، حروف نرمال."""
    base = _strip_label(text)
    base = base.replace("‌", " ").replace("_", " ").replace("-", " ")
    return normalize_text(base).lower()


REQUIRED_TEXT = {"always": "الزامی", "create": "الزامی برایِ مورد تازه", "no": "اختیاری"}

PRODUCT_COLUMNS = [
    ImportColumn("product_id", "شناسه کالا", help="شناسه‌ی داخلیِ کالا در راستی‌سی؛ برایِ شناساییِ دقیقِ کالایِ موجود (در فایلِ «خروج اطلاعات» کالاها آمده است).",
                 accepted="عددِ صحیح", example="125", blank="کالا با SKU شناسایی می‌شود؛ اگر SKU هم نباشد، کالایِ تازه ساخته می‌شود.",
                 aliases=("product id", "id", "شناسه"), width=12),
    ImportColumn("sku", "SKU", help="کدِ یکتایِ کالا در این فروشگاه؛ برایِ شناساییِ کالایِ موجود هم استفاده می‌شود.",
                 accepted="متن (بدونِ تکرار)", example="TSH-001", blank="کالایِ تازه بدونِ SKU ساخته می‌شود.", width=16, text_format=True),
    ImportColumn("name", "نام کالا", required="create", help="نامی که مشتری در فروشگاه می‌بیند.",
                 accepted="متن", example="تیشرت نخی", blank="برایِ کالایِ موجود همان نامِ فعلی می‌ماند؛ برایِ کالایِ تازه خطاست.",
                 aliases=("نام", "عنوان کالا", "name"), width=32),
    ImportColumn("status", "وضعیت", kind="choice", list_key="product_status", help="وضعیتِ نمایشِ کالا در فروشگاه.",
                 accepted="فعال / غیرفعال / پیش‌نویس", example="فعال", blank="کالایِ تازه «فعال» می‌شود؛ کالایِ موجود تغییر نمی‌کند.", width=13),
    ImportColumn("brand_code", "برند", kind="choice", list_key="brands", help="برندِ کالا؛ از فهرستِ برندهایِ همین فروشگاه انتخاب کنید.",
                 accepted="یکی از برندهایِ فهرست", example="نایک", blank="برند تغییر نمی‌کند (کالایِ تازه بدونِ برند).", width=18),
    ImportColumn("category_code", "دسته‌بندی", required="create", kind="choice", list_key="categories",
                 help="دسته‌بندیِ کالا با مسیرِ کامل (والد > فرزند). فقط زیردسته‌ها قابلِ انتخاب‌اند.",
                 accepted="یکی از مسیرهایِ فهرست", example="پوشاک > مردانه", blank="برایِ کالایِ موجود تغییر نمی‌کند؛ برایِ کالایِ تازه خطاست.",
                 aliases=("دسته", "category"), width=30),
    ImportColumn("price", "قیمت (تومان)", required="create", kind="decimal", help="قیمتِ فروشِ پایه‌ی کالا به تومان.",
                 accepted="عددِ صفر یا بزرگ‌تر", example="150000", blank="برایِ کالایِ موجود تغییر نمی‌کند؛ برایِ کالایِ تازه خطاست.",
                 aliases=("قیمت", "price"), width=16),
    ImportColumn("stock", "موجودی", kind="int", help="موجودیِ کلِ کالا. با ثبتِ یک حرکتِ انبار اعمال می‌شود و به رزروها احترام می‌گذارد.",
                 accepted="عددِ صحیحِ صفر یا بزرگ‌تر", example="40", blank="موجودی تغییر نمی‌کند.", aliases=("stock", "تعداد موجودی"), width=11),
    ImportColumn("barcode", "بارکد", help="بارکدِ کالا.", accepted="متن/عدد", example="6260123456789",
                 blank="بارکد تغییر نمی‌کند.", width=18, text_format=True),
    ImportColumn("weight_grams", "وزن (گرم)", kind="int", help="وزنِ کالا برایِ محاسبه‌ی هزینه‌ی ارسال.", accepted="عددِ صحیح", example="250",
                 blank="وزن تغییر نمی‌کند.", aliases=("وزن",), width=12),
    ImportColumn("requires_shipping", "نیاز به ارسال", kind="bool", list_key="yes_no",
                 help="آیا کالا فیزیکی است و باید ارسال شود؟", accepted="بله / خیر", example="بله",
                 blank="کالایِ تازه «بله» می‌شود؛ کالایِ موجود تغییر نمی‌کند.", width=13),
    ImportColumn("tax_class_code", "دسته مالیاتی", kind="choice", list_key="tax_classes", help="دسته‌یِ مالیاتیِ این کالا (اگر با پیش‌فرضِ فروشگاه فرق دارد).",
                 accepted="یکی از دسته‌هایِ فهرست، یا «— بدون دسته مالیاتی —» برایِ حذفِ دسته‌ی فعلی", example="عمومی",
                 blank="دسته‌یِ مالیاتی تغییر نمی‌کند.", aliases=("دسته مالیات", "مالیات"), width=18),
    ImportColumn("seo_title", "عنوان سئو", help="عنوانی که در نتایجِ گوگل نشان داده می‌شود.", accepted="متنِ کوتاه", example="خریدِ تیشرت نخی",
                 blank="تغییر نمی‌کند.", width=26),
    ImportColumn("seo_description", "توضیحات سئو", help="توضیحِ کوتاه زیرِ عنوان در نتایجِ گوگل.", accepted="متن", example="تیشرتِ نخیِ راحت و بادوام",
                 blank="تغییر نمی‌کند.", width=36),
    ImportColumn("slug", "نشانی صفحه (اسلاگ)", help="بخشی از آدرسِ صفحه‌ی کالا. معمولاً لازم نیست؛ خودکار ساخته می‌شود.", accepted="حروف، عدد و خط‌تیره",
                 example="cotton-tshirt", blank="خودکار ساخته می‌شود.", aliases=("اسلاگ", "slug"), width=22, in_template=False),
]

VARIANT_COLUMNS = [
    ImportColumn("product_id", "شناسه کالا", help="شناسه‌ی کالایِ مادر.", accepted="عددِ صحیح", example="125",
                 blank="کالا با «SKU کالا» (یا شناسه/SKU تنوع) شناسایی می‌شود.", aliases=("product id",), width=12),
    ImportColumn("product_sku", "SKU کالا", required="create", help="SKUِ کالایِ مادر؛ کالا باید از قبل وجود داشته باشد.",
                 accepted="SKUِ یک کالایِ موجود", example="TSH-001", blank="اگر شناسه‌ی کالا یا تنوع داده شده باشد، لازم نیست.", width=16, text_format=True),
    ImportColumn("variant_id", "شناسه تنوع", help="شناسه‌ی داخلیِ تنوع؛ برایِ ویرایشِ دقیقِ یک تنوعِ موجود.", accepted="عددِ صحیح", example="480",
                 blank="تنوع با SKU تنوع یا ترکیبِ ویژگی‌ها شناسایی می‌شود.", width=12),
    ImportColumn("variant_sku", "SKU تنوع", help="کدِ یکتایِ تنوع.", accepted="متن (بدونِ تکرار)", example="TSH-001-RED-L",
                 blank="SKUِ تنوع تغییر نمی‌کند.", width=18, text_format=True),
    ImportColumn("barcode", "بارکد", help="بارکدِ تنوع.", accepted="متن/عدد", example="6260123456790", blank="تغییر نمی‌کند.", width=18, text_format=True),
    ImportColumn("option_1_code", "نام ویژگی ۱", kind="choice", list_key="option_labels", help="نامِ نخستین ویژگیِ تنوع که روی کالا فعال است.",
                 accepted="یکی از ویژگی‌هایِ فعالِ همان کالا", example="رنگ", blank="اگر تنوع را با شناسه/SKU مشخص کرده‌اید، لازم نیست.", width=14),
    ImportColumn("option_1_value_code", "مقدار ویژگی ۱", help="مقدارِ همان ویژگی.", accepted="یکی از مقدارهایِ فعالِ ویژگی", example="قرمز",
                 blank="باید همراهِ نامِ ویژگی پر شود.", width=14),
    ImportColumn("option_2_code", "نام ویژگی ۲", kind="choice", list_key="option_labels", help="دومین ویژگی (اختیاری).", accepted="ویژگیِ فعالِ کالا",
                 example="سایز", blank="تنوعِ تک‌ویژگی.", width=14),
    ImportColumn("option_2_value_code", "مقدار ویژگی ۲", help="مقدارِ دومین ویژگی.", accepted="مقدارِ فعال", example="L", blank="—", width=14),
    ImportColumn("option_3_code", "نام ویژگی ۳", kind="choice", list_key="option_labels", help="سومین ویژگی (اختیاری).", accepted="ویژگیِ فعالِ کالا",
                 example="جنس", blank="—", width=14),
    ImportColumn("option_3_value_code", "مقدار ویژگی ۳", help="مقدارِ سومین ویژگی.", accepted="مقدارِ فعال", example="نخ", blank="—", width=14),
    ImportColumn("price", "تغییر قیمت (تومان)", kind="decimal", help="مبلغی که به قیمتِ پایه‌ی کالا اضافه (یا با عددِ منفی کم) می‌شود.",
                 accepted="عدد (می‌تواند منفی باشد)", example="20000", blank="تغییر نمی‌کند.", aliases=("تغییر قیمت", "قیمت"), width=16),
    ImportColumn("compare_at_price", "قیمت مقایسه‌ای (تومان)", kind="decimal", help="قیمتِ خط‌خورده‌ی قبل از تخفیف.", accepted="عدد", example="200000",
                 blank="تغییر نمی‌کند.", aliases=("قیمت مقایسه ای",), width=18),
    ImportColumn("cost", "بهای تمام‌شده (تومان)", kind="decimal", help="هزینه‌ی تمام‌شده؛ فقط برایِ شماست.", accepted="عدد", example="90000",
                 blank="تغییر نمی‌کند.", aliases=("بهای تمام شده",), width=18),
    ImportColumn("stock", "موجودی", kind="int", help="موجودیِ این تنوع (با ثبتِ حرکتِ انبار).", accepted="عددِ صحیحِ صفر یا بزرگ‌تر", example="12",
                 blank="تغییر نمی‌کند.", width=11),
    ImportColumn("weight_grams", "وزن (گرم)", kind="int", help="وزنِ تنوع.", accepted="عددِ صحیح", example="260", blank="تغییر نمی‌کند.", aliases=("وزن",), width=12),
    ImportColumn("is_active", "فعال", kind="bool", list_key="yes_no", help="آیا این تنوع قابلِ فروش است؟", accepted="بله / خیر", example="بله",
                 blank="تنوعِ تازه «بله»؛ تنوعِ موجود تغییر نمی‌کند.", width=9),
    ImportColumn("is_default", "پیش‌فرض", kind="bool", list_key="yes_no", help="تنوعی که هنگامِ ورود به صفحه‌ی کالا انتخاب شده است.",
                 accepted="بله / خیر", example="خیر", blank="«خیر».", width=10),
]

INVENTORY_COLUMNS = [
    ImportColumn("warehouse_code", "انبار", required="always", kind="choice", list_key="warehouses", help="انباری که موجودی در آن تغییر می‌کند.",
                 accepted="یکی از انبارهایِ فهرست", example="انبار مرکزی", blank="الزامی است.", width=20),
    ImportColumn("product_id", "شناسه کالا", help="شناسه‌ی داخلیِ کالا.", accepted="عددِ صحیح", example="125",
                 blank="کالا با «SKU کالا» شناسایی می‌شود.", width=12),
    ImportColumn("product_sku", "SKU کالا", help="SKUِ کالا.", accepted="SKUِ کالایِ موجود", example="TSH-001",
                 blank="یکی از «شناسه کالا» یا «SKU کالا» لازم است.", width=16, text_format=True),
    ImportColumn("variant_id", "شناسه تنوع", help="اگر موجودیِ یک تنوع را تغییر می‌دهید.", accepted="عددِ صحیح", example="480",
                 blank="موجودیِ خودِ کالا تغییر می‌کند.", width=12),
    ImportColumn("variant_sku", "SKU تنوع", help="اگر موجودیِ یک تنوع را تغییر می‌دهید.", accepted="SKUِ تنوعِ همان کالا", example="TSH-001-RED-L",
                 blank="موجودیِ خودِ کالا تغییر می‌کند.", width=18, text_format=True),
    ImportColumn("mode", "نوع عملیات", required="always", kind="choice", list_key="inventory_ops",
                 help="«تنظیم موجودی نهایی»: موجودی دقیقاً برابرِ مقدارِ واردشده می‌شود. «افزایش/کاهش موجودی»: مقدار به موجودیِ فعلی اضافه می‌شود؛ "
                      "عددِ منفی از آن کم می‌کند.",
                 accepted="تنظیم موجودی نهایی / افزایش/کاهش موجودی", example="افزایش/کاهش موجودی",
                 blank="الزامی است؛ برایِ جلوگیری از اشتباه، خالی پذیرفته نمی‌شود.", aliases=("عملیات", "حالت"), width=24),
    ImportColumn("quantity", "مقدار", required="always", kind="int",
                 help="در «تنظیم موجودی نهایی» موجودیِ نهایی؛ در «افزایش/کاهش» مقدارِ تغییر (منفی = کاهش).",
                 accepted="عددِ صحیح (در تنظیمِ نهایی نمی‌تواند منفی باشد)", example="10 یا -3", blank="الزامی است.", aliases=("تعداد",), width=11),
    ImportColumn("reason", "دلیل", help="دلیلِ تغییر (مثلاً شمارشِ سالانه).", accepted="متنِ کوتاه", example="شمارش سالانه", blank="دلیلی ثبت نمی‌شود.", width=22),
    ImportColumn("note", "توضیح", help="توضیحِ بیشتر برایِ دفترِ موجودی.", accepted="متن", example="قفسه‌ی ۲", blank="توضیحی ثبت نمی‌شود.", aliases=("یادداشت",), width=26),
]

IMPORT_SPECS = {
    ImportJob.ImportType.PRODUCTS: PRODUCT_COLUMNS,
    ImportJob.ImportType.VARIANTS: VARIANT_COLUMNS,
    ImportJob.ImportType.INVENTORY: INVENTORY_COLUMNS,
}

_NUMERIC_KEYS = {"price", "stock", "weight_grams", "quantity", "compare_at_price", "cost"}
_SEPARATORS_RE = re.compile(r"[,٬، \s]")


def alias_map(import_type: str) -> dict:
    """هدرِ نرمال‌شده ← کلیدِ داخلی. هدرهایِ فارسی، نامِ قدیمیِ داخلی (برایِ
    فایل‌هایِ CSV/اکسلِ قدیمی) و نام‌هایِ مستعار پذیرفته می‌شوند."""
    mapping = {}
    for column in IMPORT_SPECS[import_type]:
        for candidate in (column.label, column.key, *column.aliases):
            mapping.setdefault(header_key(candidate), column.key)
    return mapping


def column_by_key(import_type: str) -> dict:
    return {c.key: c for c in IMPORT_SPECS[import_type]}


# ───────────────────────────── داده‌هایِ مرجعِ Store ─────────────────────────────

@dataclass
class ReferenceData:
    store: object
    _cache: dict = field(default_factory=dict)

    def _memo(self, key, factory):
        if key not in self._cache:
            self._cache[key] = factory()
        return self._cache[key]

    @property
    def brands(self):
        return self._memo("brands", lambda: [(b.name, b.slug) for b in Brand.objects.filter(store=self.store).order_by("name")])

    @property
    def categories(self):
        return self._memo("categories", lambda: [(path, c.slug) for path, c in store_category_paths(self.store)])

    @property
    def tax_classes(self):
        return self._memo("tax", lambda: [(t.name, t.code) for t in TaxClass.objects.filter(store=self.store, is_active=True).order_by("name")])

    @property
    def all_tax_classes(self):
        return self._memo("tax_all", lambda: [(t.name, t.code) for t in TaxClass.objects.filter(store=self.store).order_by("name")])

    @property
    def warehouses(self):
        return self._memo("wh", lambda: [(w.name, w.code) for w in Warehouse.objects.filter(store=self.store).order_by("name")])

    @property
    def option_labels(self):
        def load():
            labels = ProductOption.objects.filter(product__store=self.store, is_active=True).values_list("label", flat=True).distinct()
            return sorted({label for label in labels if label})
        return self._memo("options", load)

    def lists(self, import_type: str) -> list:
        """فهرست‌هایِ مرجعِ شیتِ «فهرست‌ها» برایِ یک نوعِ واردات: ``[(کلید، عنوان، مقدارها)]``."""
        status = [label for _value, label in Product.Status.choices]
        yes_no = list(YES_NO)
        if import_type == ImportJob.ImportType.PRODUCTS:
            return [
                ("product_status", "وضعیت کالا", status),
                ("yes_no", "بله / خیر", yes_no),
                ("brands", "برندها", [n for n, _ in self.brands]),
                ("categories", "دسته‌بندی‌ها (مسیر کامل)", [p for p, _ in self.categories]),
                ("tax_classes", "دسته‌هایِ مالیاتی", [CLEAR_TAX_LABEL] + [n for n, _ in self.tax_classes]),
            ]
        if import_type == ImportJob.ImportType.VARIANTS:
            return [
                ("yes_no", "بله / خیر", yes_no),
                ("option_labels", "ویژگی‌هایِ تنوع (فعال)", self.option_labels),
            ]
        return [
            ("warehouses", "انبارها", [n for n, _ in self.warehouses]),
            ("inventory_ops", "نوع عملیات موجودی", [INV_OP_SET, INV_OP_ADJUST]),
        ]


def _name_index(pairs) -> dict:
    """نامِ نرمال‌شده ← مجموعه‌ی کدهایِ داخلی (بیش از یکی = مبهم)."""
    index: dict = {}
    for name, code in pairs:
        index.setdefault(normalization_key(name), set()).add(code)
    return index


class ValueMapper:
    """نگاشتِ مقدارهایِ قابل‌خواندنِ فارسی به کدهایِ داخلیِ یک Store. اگر مقدار
    با نامی منطبق نباشد، *همان متنِ خام* پس داده می‌شود (شاید خودِ کد را
    نوشته‌اند) و اعتبارسنجِ اصلی خطایِ «یافت نشد» را می‌دهد؛ اگر نام مبهم
    باشد، خطایِ صریح برمی‌گردد."""

    def __init__(self, store):
        self.ref = ReferenceData(store)
        self._built = {}

    def _index(self, name, factory):
        if name not in self._built:
            self._built[name] = factory()
        return self._built[name]

    def _resolve(self, index, codes, text, what):
        key = normalization_key(text)
        matches = index.get(key)
        if matches:
            if len(matches) > 1:
                return None, f"{what} «{text}» در این فروشگاه بیش از یک بار تعریف شده و مبهم است؛ نامِ آن را در پنل یکتا کنید."
            return next(iter(matches)), None
        if text in codes:
            return text, None
        return text, None

    def brand(self, text):
        index = self._index("brand", lambda: _name_index(self.ref.brands))
        return self._resolve(index, {code for _n, code in self.ref.brands}, text, "برندِ")

    def category(self, text):
        index = self._index("cat", lambda: {normalize_path_key(p): {c} for p, c in self.ref.categories})
        # مسیرِ کامل؛ در نبودِ آن، اگر «نامِ تنها» دقیقاً یک زیردسته را مشخص کند همان پذیرفته می‌شود.
        key = normalize_path_key(text)
        if key in index:
            return next(iter(index[key])), None
        by_leaf: dict = {}
        for path, code in self.ref.categories:
            by_leaf.setdefault(normalization_key(path.split(">")[-1]), set()).add(code)
        leaf = by_leaf.get(normalization_key(text))
        if leaf:
            if len(leaf) > 1:
                return None, f"نامِ دسته‌بندیِ «{text}» مبهم است؛ مسیرِ کامل را (مثلاً «والد > فرزند») از فهرست انتخاب کنید."
            return next(iter(leaf)), None
        return text, None

    def tax_class(self, text):
        if normalization_key(text) == normalization_key(CLEAR_TAX_LABEL):
            return "", None
        index = self._index("tax", lambda: _name_index(self.ref.all_tax_classes))
        return self._resolve(index, {code for _n, code in self.ref.all_tax_classes}, text, "دسته‌یِ مالیاتیِ")

    def warehouse(self, text):
        index = self._index("wh", lambda: _name_index(self.ref.warehouses))
        return self._resolve(index, {code for _n, code in self.ref.warehouses}, text, "انبارِ")

    @staticmethod
    def product_status(text):
        wanted = normalization_key(text)
        for value, label in Product.Status.choices:
            if wanted in (normalization_key(label), value):
                return value, None
        return text, None

    @staticmethod
    def inventory_op(text):
        wanted = normalization_key(text).replace("/", " ").replace("‌", " ")
        wanted = " ".join(wanted.split())
        table = {
            "تنظیم موجودی نهایی": "set_on_hand", "تنظیم موجودی": "set_on_hand", "تنظیم نهایی": "set_on_hand",
            "set on hand": "set_on_hand", "set_on_hand": "set_on_hand",
            "افزایش کاهش موجودی": "adjustment", "افزایش کاهش": "adjustment", "افزایش یا کاهش موجودی": "adjustment",
            "adjustment": "adjustment",
        }
        return table.get(wanted, text), None


# ───────────────────────────── خواندن و نگاشتِ ردیف‌ها ─────────────────────────────

@dataclass
class ParsedFile:
    rows: list  # لیستِ دیکشنری با کلیدهایِ داخلی + کلیدهایِ ``__*__``
    ignored_headers: list
    headers: list


def parse_xlsx_import(fileobj, import_type: str, store) -> ParsedFile:
    """یک XLSXِ واردات را می‌خواند و ردیف‌ها را با کلیدها/مقدارهایِ داخلی
    برمی‌گرداند — اعتبارسنج‌ها هیچ تفاوتی بینِ CSV و XLSX نمی‌بینند.

    کلیدهایِ ویژه (که اعتبارسنج‌ها نادیده می‌گیرند و ``run_import`` مصرف
    می‌کند): ``__row_number__`` (شماره‌ی ردیف در اکسل)، ``__errors__`` (ردیفِ
    رد‌شده پیش از اعتبارسنجی)، ``__warnings__``، ``__raw__`` (مقدارهایِ اصلیِ
    فایل برایِ گزارشِ خطا)."""
    headers, xrows = xlsx_utils.read_xlsx_table(fileobj)
    aliases = alias_map(import_type)
    columns = column_by_key(import_type)
    mapper = ValueMapper(store)

    keys: list = []
    ignored: list = []
    seen: set = set()
    for header in headers:
        key = aliases.get(header_key(header)) if header else None
        if header and key is None:
            ignored.append(header)
        if key is not None and key in seen:
            ignored.append(header)
            key = None
        if key is not None:
            seen.add(key)
        keys.append(key)

    if not seen:
        raise xlsx_utils.XlsxReadError(
            "ستون‌هایِ این فایل شناخته نشد. فایل را از «دانلود قالب اکسل» بگیرید و ستون‌هایِ آن را تغییر ندهید."
        )

    rows = []
    for xrow in xrows:
        data: dict = {}
        errors: list = []
        raw: list = []
        for index, key in enumerate(keys):
            text = xrow.cells[index].strip() if index < len(xrow.cells) else ""
            header = headers[index] if index < len(headers) else ""
            if header:
                raw.append((header, "(سلولِ فرمولی)" if index in xrow.formula_columns else text))
            if key is None:
                continue
            label = columns[key].base_label
            if index in xrow.formula_columns:
                errors.append(
                    f"سلولِ ستونِ «{label}» فرمول (یا خطایِ اکسل) دارد؛ فرمول‌ها پذیرفته نمی‌شوند. مقدار را به‌صورتِ عدد یا متنِ ساده وارد کنید."
                )
                continue
            if key in _NUMERIC_KEYS and text:
                text = _SEPARATORS_RE.sub("", text)
            value, problem = _map_value(import_type, key, text, mapper)
            if problem:
                errors.append(problem)
            if key == "tax_class_code" and not text:
                continue  # خالی = بدونِ تغییر (نه حذفِ دسته‌ی مالیاتی)
            data[key] = value
        if import_type == ImportJob.ImportType.INVENTORY and not data.get("mode") and not errors:
            errors.append(
                "نوعِ عملیات انتخاب نشده است: «تنظیم موجودی نهایی» یا «افزایش/کاهش موجودی» را از فهرست انتخاب کنید."
            )
        data["__row_number__"] = xrow.number
        data["__raw__"] = raw
        if errors:
            data["__errors__"] = errors
        rows.append(data)
    if rows and ignored:
        shown = "، ".join(f"«{h}»" for h in ignored[:6])
        rows[0]["__warnings__"] = [f"ستون‌هایِ {shown} شناخته نشد و نادیده گرفته شد."]
    return ParsedFile(rows=rows, ignored_headers=ignored, headers=headers)


def _map_value(import_type, key, text, mapper):
    if not text:
        return "", None
    if import_type == ImportJob.ImportType.PRODUCTS:
        if key == "brand_code":
            return mapper.brand(text)
        if key == "category_code":
            return mapper.category(text)
        if key == "tax_class_code":
            return mapper.tax_class(text)
        if key == "status":
            return mapper.product_status(text)
    elif import_type == ImportJob.ImportType.INVENTORY:
        if key == "warehouse_code":
            return mapper.warehouse(text)
        if key == "mode":
            return mapper.inventory_op(text)
    return text, None


# ───────────────────────────── پیام‌هایِ خطا (نمایشِ مشتری‌پسند) ─────────────────────────────

_FIELD_PHRASES = [
    ("تغییرِ قیمت", "price"), ("قیمتِ مقایسه", "compare_at_price"), ("بهایِ تمام", "cost"),
    ("برایِ قیمت", "price"), ("قیمت نمی‌تواند", "price"), ("برایِ موجودی", "stock"), ("موجودی نمی‌تواند", "stock"),
    ("برایِ وزن", "weight_grams"), ("برایِ تعداد", "quantity"), ("تعداد نمی‌تواند", "quantity"),
    ("وضعیتِ «", "status"), ("برندِ «", "brand_code"), ("دسته‌بندیِ «", "category_code"),
    ("دسته‌یِ مالیاتیِ", "tax_class_code"), ("انباری با کدِ", "warehouse_code"), ("حالتِ ردیفِ", "mode"),
    ("set_on_hand", "mode"), ("مقدارِ «", "option_1_value_code"), ("محورِ", "option_1_code"),
    ("شناسه‌ی کالا", "product_id"), ("شناسه‌ی تنوع", "variant_id"), ("SKUِ تنوع", "variant_sku"),
    ("SKUِ «", "sku"), ("کالایی با SKU", "product_sku"), ("تنوعی با SKU", "variant_sku"),
]
_TEXT_REPLACEMENTS = [
    ("فقط adjustment/set_on_hand", "فقط «تنظیم موجودی نهایی» یا «افزایش/کاهش موجودی»"),
    ("در حالتِ set_on_hand، تعداد نمی‌تواند منفی باشد.", "در «تنظیم موجودی نهایی»، مقدار نمی‌تواند منفی باشد."),
    ("set_on_hand", INV_OP_SET), ("adjustment", INV_OP_ADJUST),
    ("«فقط ایجاد»", f"«{MODE_LABELS[ImportJob.Mode.CREATE_ONLY]}»"),
    ("«فقط به‌روزرسانی»", f"«{MODE_LABELS[ImportJob.Mode.UPDATE_ONLY]}»"),
    ("انباری با کدِ", "انباری با نامِ"),
    ("(option_N_code/option_N_value_code)", "(نامِ ویژگی و مقدارِ ویژگی)"),
    ("هم کدِ محور و هم کدِ مقدار", "هم نامِ ویژگی و هم مقدارِ ویژگی"),
    ("محورهای", "ویژگی‌هایِ"), ("محورِ", "ویژگیِ"), ("محورها", "ویژگی‌ها"),
]


def describe_errors(import_type: str, messages) -> list:
    """پیام‌هایِ خطایِ اعتبارسنج را برایِ نمایش به مشتری بازنویسی می‌کند:
    ``[{"column": عنوانِ ستونِ فارسی یا "", "message": متنِ فارسی}]``. متنِ
    خام در پایگاه‌داده دست‌نخورده می‌ماند؛ فقط نمایش تغییر می‌کند."""
    columns = column_by_key(import_type)
    results = []
    for raw in messages or []:
        text = str(raw)
        key = None
        for match in re.finditer(r"«([a-z0-9_]+)»", text):
            if match.group(1) in columns:
                key = match.group(1)
                break
        if key is None:
            labelled = re.search(r"ستونِ «([^»]+)»", text)
            if labelled:
                key = next((k for k, c in columns.items() if c.base_label == labelled.group(1)), None)
        if key is None:
            numbered = re.search(r"option_(\d)_(?:value_)?code", text)
            if numbered:
                key = f"option_{numbered.group(1)}_code"
        if key is None:
            for phrase, candidate in _FIELD_PHRASES:
                if phrase in text and candidate in columns:
                    key = candidate
                    break
        for internal, column in columns.items():
            text = text.replace(f"«{internal}»", f"«{column.base_label}»")
        for old, new in _TEXT_REPLACEMENTS:
            text = text.replace(old, new)
        results.append({"column": columns[key].base_label if key in columns else "", "message": text})
    return results


# ───────────────────────────── قالبِ XLSX ─────────────────────────────

def _list_ranges(wb, ref: ReferenceData, import_type: str):
    """شیتِ «فهرست‌ها» را می‌نویسد و ``{list_key: (حرفِ ستون، تعداد)}`` برمی‌گرداند."""
    ws = wb.create_sheet(SHEET_LISTS)
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.tabColor = xlsx_utils.TECH_HEADER_FILL
    ranges = {}
    for index, (key, title, values) in enumerate(ref.lists(import_type), start=1):
        cell = ws.cell(row=1, column=index)
        xlsx_utils.set_text(cell, title)
        xlsx_utils.style_header_cell(cell, Column(title))
        for offset, value in enumerate(values, start=2):
            c = ws.cell(row=offset, column=index)
            xlsx_utils.set_text(c, value)
            c.font = Font(name=xlsx_utils.FONT_NAME, size=10)
            c.alignment = Alignment(horizontal="right", vertical="center")
        if not values:
            c = ws.cell(row=2, column=index)
            xlsx_utils.set_text(c, "— هنوز موردی ثبت نشده —")
            c.font = Font(name=xlsx_utils.FONT_NAME, size=10, italic=True, color=xlsx_utils.MUTED_TEXT)
        ws.column_dimensions[get_column_letter(index)].width = min(max(max([len(title)] + [len(str(v)) for v in values]) * 1.15 + 4, 16), 50)
        ranges[key] = (get_column_letter(index), len(values))
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    return ranges


def build_template_xlsx(import_type: str, store) -> bytes:
    """قالبِ راهنمادارِ XLSX را در بافتِ همین Store می‌سازد: شیتِ «داده‌ها» (فقط
    هدرِ استایل‌دار + منوهایِ کشویی)، «راهنما» و «فهرست‌ها»."""
    specs = [c for c in IMPORT_SPECS[import_type] if c.in_template]
    type_label = dict(ImportJob.ImportType.choices)[import_type]
    ref = ReferenceData(store)
    wb = xlsx_utils.new_workbook(f"قالبِ ورودِ اطلاعات {type_label}")
    data_ws = wb.active
    data_ws.title = SHEET_DATA

    table_columns = [
        Column(
            f"{c.label} *" if c.required != "no" else c.label, width=c.width, required=(c.required != "no"),
            note=f"{REQUIRED_TEXT[c.required]}. {c.help}", text_format=c.text_format,
        )
        for c in specs
    ]
    xlsx_utils.write_table_sheet(data_ws, table_columns, [], header_height=38)

    guide_ws = wb.create_sheet(SHEET_GUIDE)
    ranges = _list_ranges(wb, ref, import_type)

    last_row = TEMPLATE_VALIDATION_ROWS + 1
    for index, column in enumerate(specs, start=1):
        letter = get_column_letter(index)
        span = f"{letter}2:{letter}{last_row}"
        if column.text_format:
            for row in range(2, 502):
                data_ws.cell(row=row, column=index).number_format = "@"
        validation = None
        if column.list_key and column.list_key in ranges and ranges[column.list_key][1] > 0:
            list_letter, count = ranges[column.list_key]
            validation = DataValidation(
                type="list", formula1=f"'{SHEET_LISTS}'!${list_letter}$2:${list_letter}${count + 1}",
                allow_blank=True, showErrorMessage=True, errorStyle="stop",
                errorTitle="مقدارِ نامعتبر", error="لطفاً یکی از گزینه‌هایِ فهرست را انتخاب کنید.",
            )
        elif column.kind == "int" and import_type != ImportJob.ImportType.INVENTORY:
            validation = DataValidation(
                type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                showErrorMessage=True, errorTitle="عددِ نامعتبر", error="یک عددِ صحیحِ صفر یا بزرگ‌تر وارد کنید.",
            )
        elif column.kind == "int":
            validation = DataValidation(
                type="whole", operator="between", formula1="-1000000000", formula2="1000000000", allow_blank=True,
                showErrorMessage=True, errorTitle="عددِ نامعتبر", error="یک عددِ صحیح وارد کنید (عددِ منفی = کاهش).",
            )
        elif column.kind == "decimal" and not (import_type == ImportJob.ImportType.VARIANTS and column.key == "price"):
            validation = DataValidation(
                type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                showErrorMessage=True, errorTitle="عددِ نامعتبر", error="یک عددِ صفر یا بزرگ‌تر وارد کنید.",
            )
        elif column.kind == "decimal":
            validation = DataValidation(
                type="decimal", operator="between", formula1="-1000000000", formula2="1000000000", allow_blank=True,
                showErrorMessage=True, errorTitle="عددِ نامعتبر", error="یک عدد وارد کنید (می‌تواند منفی باشد).",
            )
        if validation is not None:
            validation.add(span)
            data_ws.add_data_validation(validation)

    intro = [
        "۱) این شیتِ «داده‌ها» را پر کنید؛ هر ردیف یک مورد است. نخستین ردیف (عنوانِ ستون‌ها) را تغییر ندهید و ستونی را حذف یا جابه‌جا نکنید.",
        "۲) عنوانِ قهوه‌ایِ دارایِ ستاره (*) یعنی الزامی؛ عنوانِ سبزِ تیره یعنی اختیاری. جزئیاتِ هر ستون در جدولِ پایین آمده است.",
        "۳) جایی که منویِ کشویی دارد، مقدار را از فهرست انتخاب کنید (فهرست‌ها از اطلاعاتِ همین فروشگاه ساخته شده‌اند).",
        "۴) فایل را با قالبِ «Excel Workbook (*.xlsx)» ذخیره و در بخشِ «ورود اطلاعات» بارگذاری کنید. فایل‌هایِ ماکرودار (xlsm) پذیرفته نمی‌شوند.",
        "۵) فرمول‌هایِ اکسل اجرا نمی‌شوند؛ اگر سلولی فرمول داشته باشد آن ردیف خطا دارد و کلِ فایل اجرا نمی‌شود. مقدار را به‌صورتِ عدد یا متنِ ساده وارد کنید.",
        "۶) ارقامِ فارسی و جداکننده‌یِ هزارگان در عددها مشکلی ایجاد نمی‌کنند. هر فایل حداکثر ۲۰٬۰۰۰ ردیف دارد.",
        "بارگذاری فایل هیچ تغییری ایجاد نمی‌کند؛ ابتدا همه ردیف‌ها بررسی می‌شوند. اگر حتی یک ردیف خطا داشته باشد فایل اجرا نمی‌شود و باید آن را اصلاح و دوباره بارگذاری کنید. فقط وقتی همه ردیف‌ها معتبر باشند و شما تأیید کنید، اطلاعاتِ فروشگاه تغییر می‌کند (همه یا هیچ).",
    ]
    if import_type == ImportJob.ImportType.INVENTORY:
        intro.insert(0, (
            "نوع عملیات: «تنظیم موجودی نهایی» یعنی موجودیِ کالا دقیقاً برابرِ «مقدار» شود (مثلاً مقدار ۱۰ ⇒ موجودی می‌شود ۱۰). "
            "«افزایش/کاهش موجودی» یعنی «مقدار» به موجودیِ فعلی اضافه شود؛ عددِ منفی از موجودی کم می‌کند (مثلاً ‎-۳‎ ⇒ سه عدد کم می‌شود). "
            "تغییری که موجودیِ قابلِ فروش را زیرِ رزروهایِ فعال ببرد رد می‌شود."
        ))
    if import_type == ImportJob.ImportType.PRODUCTS:
        intro.insert(0, (
            "هر ردیف یا یک کالایِ تازه است یا بروزرسانیِ یک کالایِ موجود. کالایِ موجود با «شناسه کالا» (در اولویت) یا «SKU» شناخته می‌شود؛ "
            "در ردیفِ بروزرسانی فقط ستون‌هایی را پر کنید که می‌خواهید عوض شوند، بقیه خالی بمانند."
        ))

    rows = [
        (c.label, REQUIRED_TEXT[c.required], c.help, c.accepted, c.example, c.blank)
        for c in specs
    ]
    xlsx_utils.write_info_sheet(
        guide_ws, title=f"راهنمایِ ورودِ اطلاعات {type_label}",
        facts=[("فروشگاه", store.name), ("تاریخ تهیه‌ی قالب", timezone.now())],
        paragraphs=intro,
        table_header=("ستون", "الزامی؟", "توضیح", "مقدارهایِ مجاز", "مثال", "اگر خالی بماند"),
        table_rows=rows, widths=(22, 18, 44, 34, 22, 40),
    )
    return xlsx_utils.workbook_to_bytes(wb)


# ───────────────────────────── گزارشِ خطا ─────────────────────────────

ROW_STATUS_LABELS = {"invalid": "نامعتبر", "failed": "ناموفق"}


def build_error_report_xlsx(job, row_results, source_rows: dict, import_type: str) -> bytes:
    """گزارشِ خطایِ XLSX: برایِ هر ردیفِ نامعتبر/ناموفق، شماره‌ی ردیف، وضعیت، خطاها،
    هشدارها و مقدارهایِ اصلیِ همان ردیف؛ سلولِ مرتبط با خطا قرمز می‌شود."""
    columns_by_key = column_by_key(import_type)
    wb = xlsx_utils.new_workbook(f"گزارشِ خطایِ ورودِ اطلاعات {job.pk}")
    ws = wb.active
    ws.title = "خطاها"

    header_pool: list = []
    for row_number in source_rows:
        for header, _value in source_rows[row_number].get("__raw__", []):
            if header not in header_pool:
                header_pool.append(header)

    alias = alias_map(import_type)
    fixed = [
        Column("ردیف در اکسل", "id", width=11), Column("وضعیت", width=11),
        Column("خطاها", width=60, wrap=True), Column("هشدارها", width=40, wrap=True),
    ]
    table_columns = fixed + [Column(h, width=18, wrap=True, tier="tech") for h in header_pool]

    results = list(row_results)
    described = []
    table_rows = []
    for result in results:
        details = describe_errors(import_type, result.errors)
        described.append(details)
        raw = dict(source_rows.get(result.row_number, {}).get("__raw__", []))
        table_rows.append([
            result.row_number, ROW_STATUS_LABELS.get(result.status, result.status),
            "\n".join(f"• {d['message']}" for d in details),
            "\n".join(f"• {w}" for w in result.warnings),
            *[raw.get(h, "") for h in header_pool],
        ])
    xlsx_utils.write_table_sheet(ws, table_columns, table_rows, empty_message="ردیفِ دارایِ خطا وجود ندارد.")

    error_fill = PatternFill("solid", start_color=xlsx_utils.ERROR_FILL, end_color=xlsx_utils.ERROR_FILL)
    for offset, (result, details) in enumerate(zip(results, described), start=2):
        for column_index in (2, 3):
            cell = ws.cell(row=offset, column=column_index)
            cell.fill = error_fill
            cell.font = Font(name=xlsx_utils.FONT_NAME, size=10, bold=(column_index == 2), color=xlsx_utils.ERROR_TEXT)
        flagged = {d["column"] for d in details if d["column"]}
        for position, header in enumerate(header_pool, start=len(fixed) + 1):
            key = alias.get(header_key(header))
            if key and columns_by_key[key].base_label in flagged:
                cell = ws.cell(row=offset, column=position)
                cell.fill = error_fill
                cell.font = Font(name=xlsx_utils.FONT_NAME, size=10, bold=True, color=xlsx_utils.ERROR_TEXT)

    guide = wb.create_sheet(SHEET_GUIDE)
    xlsx_utils.write_info_sheet(
        guide, title="گزارشِ خطایِ ورودِ اطلاعات",
        facts=[
            ("نوع اطلاعات", job.get_import_type_display()),
            ("نامِ فایل", job.original_filename or "—"),
            ("ردیف‌هایِ دارایِ خطا", len(results)),
        ],
        paragraphs=[
            "تا وقتی حتی یک ردیف خطا دارد، کلِ فایل اجرا نمی‌شود و هیچ تغییری در فروشگاه اعمال نشده است. سلول‌هایِ قرمز نشان می‌دهند خطا از کدام ستون است.",
            "همه‌ی خطاها را در همین فایلِ اصلی اصلاح کنید و آن را دوباره در «ورود اطلاعات» بارگذاری کنید.",
        ],
    )
    return xlsx_utils.workbook_to_bytes(wb)
