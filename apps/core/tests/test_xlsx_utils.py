"""Unit tests for the shared XLSX layer (``apps.core.services.xlsx_utils``):
literal-text safety, sheet formatting, upload validation (extension, size, ZIP
structure, macros, ZIP bombs) and the bounded, formula-rejecting reader."""

import io
import zipfile
from datetime import date, datetime, timezone as dt_timezone
from decimal import Decimal
from types import SimpleNamespace
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from apps.core.services import xlsx_utils
from apps.core.services.xlsx_utils import Column, XlsxUploadError, safe_download_filename

XLSX_MIME = xlsx_utils.XLSX_CONTENT_TYPE


def _book(rows, *, headers=("الف", "ب"), title="داده‌ها"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _upload(content, name="f.xlsx", content_type=XLSX_MIME):
    return SimpleUploadedFile(name, content, content_type=content_type)


class LiteralTextTests(SimpleTestCase):
    def test_formula_prefixes_become_quote_prefixed_literal_strings(self):
        wb = Workbook()
        ws = wb.active
        values = ["=1+1", "+1", "-1", "@A1", "\tTab", "=HYPERLINK(\"u\")", "عادی", "a=b"]
        for index, value in enumerate(values, start=1):
            xlsx_utils.set_text(ws.cell(row=index, column=1), value)
        buffer = io.BytesIO()
        wb.save(buffer)
        reopened = load_workbook(io.BytesIO(buffer.getvalue())).active
        for index, value in enumerate(values, start=1):
            cell = reopened.cell(row=index, column=1)
            self.assertEqual(cell.value, value)  # exact round trip, no apostrophe added
            self.assertEqual(cell.data_type, "s")
            self.assertEqual(bool(cell.quotePrefix), value.startswith(("=", "+", "-", "@", "\t")), value)
        xml = zipfile.ZipFile(io.BytesIO(buffer.getvalue())).read("xl/worksheets/sheet1.xml").decode()
        self.assertNotIn("<f>", xml)

    def test_illegal_xml_characters_are_removed_and_long_text_is_capped(self):
        self.assertEqual(xlsx_utils.clean_text("a\x00b\x07c\x1fd"), "abcd")
        self.assertEqual(len(xlsx_utils.clean_text("x" * 40000)), xlsx_utils.EXCEL_MAX_CELL_CHARS)

    def test_set_value_types(self):
        wb = Workbook()
        ws = wb.active
        cases = [
            (Decimal("150000.00"), "money", 150000, "#,##0"),
            (Decimal("12.5"), "decimal", 12.5, "#,##0.##"),
            ("42", "int", 42, "#,##0"),
            (7.0, "int", 7, "#,##0"),
            (True, "text", "بله", "General"),
        ]
        for index, (value, kind, expected, number_format) in enumerate(cases, start=1):
            cell = ws.cell(row=index, column=1)
            xlsx_utils.set_value(cell, value, kind)
            self.assertEqual(cell.value, expected)
            self.assertEqual(cell.number_format, number_format)
        stamp = ws.cell(row=10, column=1)
        xlsx_utils.set_value(stamp, datetime(2026, 1, 1, 8, 30, tzinfo=dt_timezone.utc), "datetime")
        self.assertIsNone(stamp.value.tzinfo)
        self.assertEqual((stamp.value.hour, stamp.value.minute), (12, 0))  # Asia/Tehran = UTC+3:30
        day = ws.cell(row=11, column=1)
        xlsx_utils.set_value(day, date(2026, 3, 1), "date")
        self.assertEqual(day.number_format, "yyyy/mm/dd")
        empty = ws.cell(row=12, column=1)
        xlsx_utils.set_value(empty, None, "money")
        self.assertIsNone(empty.value)


class TableSheetTests(SimpleTestCase):
    def _sheet(self, rows, empty_message=""):
        wb = xlsx_utils.new_workbook("t")
        ws = wb.active
        ws.title = "نمونه"
        columns = [Column("نام", width=20), Column("قیمت", "money"), Column("شناسه", "id", tier="tech")]
        count = xlsx_utils.write_table_sheet(ws, columns, rows, empty_message=empty_message)
        reopened = load_workbook(io.BytesIO(xlsx_utils.workbook_to_bytes(wb)))["نمونه"]
        return count, reopened

    def test_header_freeze_filter_rtl_and_widths(self):
        count, ws = self._sheet([["الف", 1000, 1], ["ب", 2500, 2]])
        self.assertEqual(count, 2)
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:C3")
        self.assertEqual(ws.column_dimensions["A"].width, 20)
        self.assertEqual(ws["B2"].number_format, "#,##0")
        self.assertEqual(ws["B3"].value, 2500)
        self.assertTrue(ws["A1"].font.bold)
        self.assertNotEqual(ws["C1"].fill.start_color.rgb, ws["A1"].fill.start_color.rgb)

    def test_empty_sheet_gets_a_visible_explanation(self):
        count, ws = self._sheet([], empty_message="موردی نیست")
        self.assertEqual(count, 0)
        self.assertEqual(ws["A2"].value, "موردی نیست")
        self.assertIn("A2:C2", [str(r) for r in ws.merged_cells.ranges])

    def test_header_cells_are_text_even_if_they_look_like_formulas(self):
        wb = xlsx_utils.new_workbook("t")
        ws = wb.active
        xlsx_utils.write_table_sheet(ws, [Column("=SUM(A1)")], [["x"]])
        reopened = load_workbook(io.BytesIO(xlsx_utils.workbook_to_bytes(wb))).active
        self.assertEqual(reopened["A1"].data_type, "s")


class UploadValidationTests(SimpleTestCase):
    def test_valid_xlsx_and_csv_are_accepted(self):
        self.assertEqual(xlsx_utils.validate_import_upload(_upload(_book([]))), ".xlsx")
        self.assertEqual(
            xlsx_utils.validate_import_upload(_upload(b"a,b\n1,2\n", "f.csv", "text/csv")), ".csv",
        )
        # browsers often send octet-stream
        self.assertEqual(
            xlsx_utils.validate_import_upload(_upload(_book([]), content_type="application/octet-stream")), ".xlsx",
        )

    def test_dangerous_or_unsupported_extensions_are_rejected(self):
        for name in ("a.xlsm", "a.XLSM", "a.xlsb", "a.xls", "a.xltm", "a.ods", "a.exe", "a.zip", "noext", "a.xlsx.exe"):
            with self.subTest(name=name), self.assertRaises(XlsxUploadError):
                xlsx_utils.validate_import_upload(_upload(_book([]), name))

    def test_paths_and_null_bytes_in_names_are_rejected(self):
        # (Django's UploadedFile already strips directories from real uploads, so the
        # guard is exercised with a plain object to prove it holds independently.)
        for name in ("../a.xlsx", "a/b.xlsx", "a\\b.xlsx", "a\x00.xlsx"):
            fake = SimpleNamespace(name=name, size=10, content_type=XLSX_MIME)
            with self.subTest(name=name), self.assertRaises(XlsxUploadError):
                xlsx_utils.validate_import_upload(fake)

    def test_non_zip_and_corrupt_files_are_rejected(self):
        for content in (b"", b"plain text", b"PK\x03\x04garbage", _book([])[:200]):
            with self.subTest(content=content[:8]), self.assertRaises(XlsxUploadError):
                xlsx_utils.validate_import_upload(_upload(content))

    def test_zip_without_workbook_parts_is_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as z:
            z.writestr("hello.txt", "hi")
        with self.assertRaises(XlsxUploadError):
            xlsx_utils.validate_import_upload(_upload(buffer.getvalue()))

    def test_macro_enabled_content_is_rejected(self):
        original = zipfile.ZipFile(io.BytesIO(_book([])))
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as z:
            for item in original.infolist():
                z.writestr(item, original.read(item.filename))
            z.writestr("xl/vbaProject.bin", b"x")
        with self.assertRaisesMessage(XlsxUploadError, "ماکرو"):
            xlsx_utils.validate_import_upload(_upload(buffer.getvalue()))
        # macro-enabled content type declared in [Content_Types].xml
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as z:
            for item in original.infolist():
                data = original.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml",
                    )
                z.writestr(item, data)
        with self.assertRaises(XlsxUploadError):
            xlsx_utils.validate_import_upload(_upload(buffer.getvalue()))

    def test_size_content_type_and_zip_bomb_limits(self):
        big = SimpleNamespace(name="a.xlsx", size=xlsx_utils.MAX_IMPORT_FILE_SIZE_BYTES + 1, content_type=XLSX_MIME)
        with self.assertRaises(XlsxUploadError):
            xlsx_utils.validate_import_upload(big)
        with self.assertRaises(XlsxUploadError):
            xlsx_utils.validate_import_upload(_upload(_book([]), content_type="application/x-msdownload"))
        with mock.patch.object(xlsx_utils, "MAX_UNCOMPRESSED_BYTES", 100):
            with self.assertRaises(XlsxUploadError):
                xlsx_utils.validate_import_upload(_upload(_book([])))
        with mock.patch.object(xlsx_utils, "MAX_ZIP_ENTRIES", 2):
            with self.assertRaises(XlsxUploadError):
                xlsx_utils.validate_import_upload(_upload(_book([])))


class ReaderTests(SimpleTestCase):
    def _read(self, content, **kwargs):
        return xlsx_utils.read_xlsx_table(io.BytesIO(content), **kwargs)

    def test_reads_headers_rows_numbers_and_normalises_values(self):
        headers, rows = self._read(_book([
            ["۱۲۳", 5], [12.0, 7.5], [True, None], [datetime(2026, 1, 2), "y"],
        ]))
        self.assertEqual(headers, ["الف", "ب"])
        self.assertEqual([r.number for r in rows], [2, 3, 4, 5])
        self.assertEqual(rows[0].cells, ["۱۲۳", "5"])
        self.assertEqual(rows[1].cells, ["12", "7.5"])
        self.assertEqual(rows[2].cells, ["true", ""])
        self.assertEqual(rows[3].cells, ["2026-01-02T00:00:00", "y"])
        self.assertEqual(xlsx_utils.cell_to_text("x\x00y\x07"), "xy")

    def test_formula_cells_are_flagged_never_evaluated(self):
        _headers, rows = self._read(_book([["=1+1", "ok"], ["x", "=SUM(A1:A2)"]]))
        self.assertEqual(rows[0].formula_columns, [0])
        self.assertEqual(rows[0].cells, ["", "ok"])  # neither the formula text nor a computed value
        self.assertEqual(rows[1].formula_columns, [1])

    def test_blank_rows_skipped_and_row_limit_enforced(self):
        _h, rows = self._read(_book([["a", "b"], [None, None], ["c", "d"]]))
        self.assertEqual([r.number for r in rows], [2, 4])
        with self.assertRaises(xlsx_utils.XlsxRowLimitExceededError):
            self._read(_book([["a", "b"]] * 5), max_rows=3)

    def test_long_cells_are_truncated_to_the_field_limit(self):
        _h, rows = self._read(_book([["x" * 5000, "y"]]))
        self.assertEqual(len(rows[0].cells[0]), xlsx_utils.MAX_FIELD_LENGTH)

    def test_prefers_data_sheet_and_skips_guide(self):
        wb = Workbook()
        guide = wb.active
        guide.title = "راهنما"
        guide.append(["راهنما"])
        data = wb.create_sheet("داده‌ها")
        data.append(["الف"])
        data.append(["مقدار"])
        buffer = io.BytesIO()
        wb.save(buffer)
        headers, rows = self._read(buffer.getvalue())
        self.assertEqual((headers, rows[0].cells), (["الف"], ["مقدار"]))

    def test_corrupt_workbook_raises_a_readable_error(self):
        with self.assertRaises(xlsx_utils.XlsxReadError):
            self._read(b"not a workbook")


class SafeDownloadFilenameTests(SimpleTestCase):
    def test_keeps_persian_and_english_basenames(self):
        self.assertEqual(safe_download_filename("محصولات پاییز.xlsx", "xlsx", "f"), "محصولات پاییز.xlsx")
        self.assertEqual(safe_download_filename("Spring v2.XLSX", "xlsx", "f"), "Spring v2.xlsx")
        self.assertEqual(safe_download_filename("old.csv", "csv", "f"), "old.csv")

    def test_adds_the_real_extension_when_missing_or_different(self):
        self.assertEqual(safe_download_filename("catalog", "xlsx", "f"), "catalog.xlsx")
        self.assertEqual(safe_download_filename("catalog.v2", "xlsx", "f"), "catalog.v2.xlsx")

    def test_drops_directories_control_chars_and_reserved_characters(self):
        self.assertEqual(safe_download_filename("../../etc/passwd.xlsx", "xlsx", "f"), "passwd.xlsx")
        self.assertEqual(safe_download_filename("C:\\temp\\a.xlsx", "xlsx", "f"), "a.xlsx")
        self.assertEqual(safe_download_filename("a\r\nb\x00c.xlsx", "xlsx", "f"), "abc.xlsx")
        self.assertEqual(safe_download_filename('x"y;z<>.xlsx', "xlsx", "f"), "xyz.xlsx")
        self.assertEqual(safe_download_filename("\u202egpj.xlsx", "xlsx", "f"), "gpj.xlsx")

    def test_falls_back_and_limits_length(self):
        for bad in ("", "..", "...xlsx", " . ", "/"):
            self.assertEqual(safe_download_filename(bad, "xlsx", "fallback-1"), "fallback-1.xlsx", bad)
        self.assertLessEqual(len(safe_download_filename("a" * 500 + ".xlsx", "xlsx", "f")), 125)
