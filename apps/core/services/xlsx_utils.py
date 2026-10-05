"""ابزارهای مشترکِ XLSX برای صادرات/واردات — جایگزینِ ``csv_utils`` برای
فایل‌های اکسل (ADR-51/ADR-62 به‌علاوه‌ی سیاستِ XLSX در
``docs/design/XLSX_IMPORT_EXPORT_V1.md``).

هیچ ویو یا سرویسی نباید مستقیماً ``openpyxl`` را برایِ تولید/خواندنِ فایلِ
کاربر صدا بزند؛ همه از این ماژول عبور می‌کنند تا سه تضمینِ امنیتی/کیفی
همیشه با هم برقرار باشند:

* **هرگز فرمول تولید نمی‌شود.** هر رشته‌ی کاربر صریحاً به‌صورتِ سلولِ
  *رشته‌ایِ متنی* نوشته می‌شود (``data_type = "s"``)، و اگر با ``= + - @`` یا
  کاراکترِ کنترلی شروع شود ویژگیِ «quote prefix» اکسل هم روشن می‌شود تا
  حتی با ویرایشِ بعدیِ سلول به فرمول تبدیل نشود. مقدار دقیقاً همان رشته‌ی
  اصلی می‌ماند (برخلافِ CSV که آپاستروف اضافه می‌کرد).
* **هیچ فایلِ ماکرودار پذیرفته نمی‌شود** (``.xlsm``/``.xlsb``/``.xls``/...)،
  و ساختارِ ZIPِ فایل (حجمِ باز‌شده، تعدادِ اجزاء، ``vbaProject.bin``) پیش
  از هر پردازشی بررسی می‌شود.
* **سلول‌هایِ فرمولی در واردات اجرا نمی‌شوند** — خواننده آن‌ها را شناسایی
  و به‌صورتِ خطایِ ردیف گزارش می‌کند (نه مقدارِ محاسبه‌شده، نه متنِ فرمول).
"""

from __future__ import annotations

import io
import math
import re
import zipfile
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.services.csv_utils import MAX_FIELD_LENGTH, MAX_IMPORT_FILE_SIZE_BYTES, MAX_IMPORT_ROWS

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CONTENT_TYPE = "text/csv"

ALLOWED_UPLOAD_CONTENT_TYPES = frozenset({
    XLSX_CONTENT_TYPE, "application/octet-stream", "application/zip", "application/x-zip-compressed",
    "text/csv", "application/vnd.ms-excel", "text/plain",
})
#: پسوندهایی که عمداً رد می‌شوند (ماکرودار یا قالبِ قدیمیِ باینری).
REJECTED_UPLOAD_EXTENSIONS = (".xlsm", ".xlsb", ".xls", ".xltm", ".xltx", ".xlam", ".ods")

MAX_ZIP_ENTRIES = 200
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_COLUMNS = 60
MAX_TRAILING_BLANK_ROWS = 2_000
EXCEL_MAX_CELL_CHARS = 32_000

# ───────────────────────────── ظاهر (Style tokens) ─────────────────────────────
FONT_NAME = "Tahoma"  # روی ویندوز/مک موجود است و حروفِ فارسی را درست نشان می‌دهد
BRAND = "214F40"
BRAND_SOFT = "EDF1E8"
ACCENT = "DCEDA6"
MUTED_TEXT = "6B746F"
TECH_HEADER_FILL = "D9DDD6"
REQUIRED_HEADER_FILL = "8E632F"
LINE = "D5DBD1"
ZEBRA = "F6F8F3"
ERROR_FILL = "FBE3E3"
ERROR_TEXT = "9D3434"
WARNING_FILL = "FFF1CC"
OK_FILL = "E3F1DC"

_THIN = Side(style="thin", color=LINE)
CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")


class XlsxUploadError(Exception):
    """فایلِ آپلودشده رد شد — پیام برایِ نمایشِ مستقیم به کاربر امن است."""


@dataclass
class Column:
    """یک ستونِ شیتِ داده. ``kind``: text/int/money/decimal/date/datetime/id."""

    header: str
    kind: str = "text"
    width: float | None = None
    tier: str = "main"  # main | tech (فنی، کم‌رنگ و انتهایی)
    note: str = ""  # توضیحِ کوتاه: کامنتِ هدر + ردیفِ راهنما
    wrap: bool = False
    required: bool = False
    text_format: bool = False  # فرمتِ «متن» (مثلاً بارکد/SKU تا صفرِ ابتدایی حذف نشود)


_NUMBER_FORMATS = {
    "int": "#,##0",
    "money": "#,##0",
    "decimal": "#,##0.##",
    "id": "0",
    "date": "yyyy/mm/dd",
    "datetime": "yyyy/mm/dd hh:mm",
}


# ───────────────────────────── مقدارها ─────────────────────────────

def clean_text(value) -> str:
    """رشته‌ی امن برایِ نوشتن: بدونِ کاراکترهایِ غیرمجازِ XML و با سقفِ طول."""
    text = _ILLEGAL_XML_RE.sub("", str(value))
    if len(text) > EXCEL_MAX_CELL_CHARS:
        text = text[:EXCEL_MAX_CELL_CHARS]
    return text


def set_text(cell, value) -> None:
    """مقدارِ کاربر را به‌صورتِ *متنِ لفظی* می‌نویسد — هرگز فرمول نمی‌شود."""
    text = clean_text(value)
    cell.value = text
    cell.data_type = "s"
    if text.startswith(_FORMULA_PREFIXES):
        cell.quotePrefix = True


def _naive_local(value: datetime) -> datetime:
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.replace(tzinfo=None, microsecond=0)


def set_value(cell, value, kind: str = "text", *, with_format: bool = True) -> None:
    """مقدار را با نوعِ مناسب می‌نویسد: عددِ واقعی، تاریخِ واقعی یا متنِ لفظی."""
    if value is None or value == "":
        cell.value = None
        return
    if kind in ("int", "money", "decimal", "id") and not isinstance(value, bool):
        number = value
        if isinstance(value, str):
            try:
                number = Decimal(value)
            except InvalidOperation:
                set_text(cell, value)
                return
        if isinstance(number, Decimal):
            number = int(number) if number == number.to_integral_value() else float(number)
        elif isinstance(number, float) and number.is_integer():
            number = int(number)
        cell.value = number
        if with_format:
            cell.number_format = _NUMBER_FORMATS[kind]
        return
    if isinstance(value, datetime):
        cell.value = _naive_local(value)
        if with_format:
            cell.number_format = _NUMBER_FORMATS["datetime"]
        return
    if isinstance(value, date):
        cell.value = value
        if with_format:
            cell.number_format = _NUMBER_FORMATS["date"]
        return
    if isinstance(value, bool):
        set_text(cell, "بله" if value else "خیر")
        return
    set_text(cell, value)


# ───────────────────────────── نوشتنِ شیت‌ها ─────────────────────────────

def _font(**kwargs) -> Font:
    return Font(name=FONT_NAME, size=kwargs.pop("size", 10), **kwargs)


def style_header_cell(cell, column: Column) -> None:
    if column.required:
        fill, color = REQUIRED_HEADER_FILL, "FFFFFF"
    elif column.tier == "tech":
        fill, color = TECH_HEADER_FILL, MUTED_TEXT
    else:
        fill, color = BRAND, "FFFFFF"
    cell.font = _font(bold=True, color=color)
    cell.fill = PatternFill("solid", start_color=fill, end_color=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = CELL_BORDER
    if column.note:
        comment = Comment(clean_text(column.note), "راستی‌سی")
        comment.width, comment.height = 260, 90
        cell.comment = comment


def _estimate_width(column: Column, sample: list) -> float:
    if column.width:
        return column.width
    longest = max([len(column.header)] + [len(str(v)) for v in sample if v not in (None, "")])
    base = {"date": 12, "datetime": 17, "money": 14, "int": 10, "decimal": 12, "id": 10}.get(column.kind)
    width = base or min(max(longest * 1.15 + 3, 10), 48)
    return max(width, min(len(column.header) * 1.15 + 3, 30))


def write_table_sheet(ws, columns: list[Column], rows, *, empty_message: str = "", header_height: float = 34) -> int:
    """شیتِ داده را می‌نویسد: هدرِ استایل‌دار، ردیف‌هایِ بدونِ فرمول، فریزِ هدر،
    فیلترِ خودکار، عرضِ ستون، RTL. تعدادِ ردیف‌هایِ داده را برمی‌گرداند."""
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = BRAND
    ws.sheet_format.defaultRowHeight = 20

    for index, column in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index)
        set_text(cell, column.header)
        style_header_cell(cell, column)
    ws.row_dimensions[1].height = header_height

    # هش‌کردنِ شیءهایِ استایلِ openpyxl برایِ هر سلول گران است (صادراتِ ۲۰٬۰۰۰ ردیفی را
    # چند ده ثانیه کند می‌کرد). به‌جایش برایِ هر (ستون، ردیفِ زوج/فرد) *یک بار* استایل را
    # با API معمولی روی یک سلول می‌گذاریم و بقیه‌ی سلول‌ها آرایه‌ی استایلِ آن را کپی می‌کنند.
    zebra_fill = PatternFill("solid", start_color=ZEBRA, end_color=ZEBRA)
    prototypes: dict = {}

    def style_cell(cell, column_index: int, column: Column, zebra: bool) -> None:
        key = (column_index, zebra)
        proto = prototypes.get(key)
        if proto is not None:
            cell._style = copy(proto)
            return
        is_number = column.kind in ("int", "money", "decimal", "id")
        cell.font = _font(color=MUTED_TEXT if column.tier == "tech" else "1F2D29")
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=column.wrap or (not is_number))
        cell.border = CELL_BORDER
        if column.text_format:
            cell.number_format = "@"
        elif column.kind in _NUMBER_FORMATS:
            cell.number_format = _NUMBER_FORMATS[column.kind]
        if zebra:
            cell.fill = zebra_fill
        prototypes[key] = copy(cell._style)

    sample_values: list[list] = [[] for _ in columns]
    count = 0
    for row in rows:
        count += 1
        excel_row = count + 1
        zebra = count % 2 == 0
        tallest = 1
        for index, (column, value) in enumerate(zip(columns, row), start=1):
            cell = ws.cell(row=excel_row, column=index)
            set_value(cell, value, column.kind, with_format=False)
            style_cell(cell, index, column, zebra)
            if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_PREFIXES):
                cell.quotePrefix = True  # پس از کپیِ استایل (که پرچم را بازنشانی می‌کند)
            if len(sample_values[index - 1]) < 200 and value not in (None, ""):
                sample_values[index - 1].append(value)
            if column.wrap and isinstance(value, str) and column.width:
                tallest = max(tallest, math.ceil(len(value) / max(column.width - 2, 8)))
        ws.row_dimensions[excel_row].height = min(max(22, 15 * tallest + 6), 130)

    if count == 0 and empty_message:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(columns), 1))
        cell = ws.cell(row=2, column=1)
        set_text(cell, empty_message)
        cell.font = _font(italic=True, bold=True, color=REQUIRED_HEADER_FILL, size=11)
        cell.fill = PatternFill("solid", start_color=WARNING_FILL, end_color=WARNING_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 42

    for index, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = _estimate_width(column, sample_values[index - 1])

    last_row = max(count + 1, 2 if (count == 0 and empty_message) else 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    return count


def write_info_sheet(ws, *, title: str, facts: list[tuple[str, object]], paragraphs: list[str],
                     table_header: tuple[str, ...] = (), table_rows: list[tuple] | None = None,
                     widths: tuple[float, ...] = (28, 70)) -> None:
    """شیتِ «راهنما»: عنوان، چند واقعیتِ کلیدی-مقدار، توضیحات، و یک جدولِ اختیاری."""
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ACCENT
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    span = max(len(widths), 2)

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1)
    set_text(cell, title)
    cell.font = _font(bold=True, size=15, color="FFFFFF")
    cell.fill = PatternFill("solid", start_color=BRAND, end_color=BRAND)
    cell.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws.row_dimensions[row].height = 38
    row += 2

    for label, value in facts:
        key = ws.cell(row=row, column=1)
        set_text(key, label)
        key.font = _font(bold=True, color=BRAND)
        key.fill = PatternFill("solid", start_color=BRAND_SOFT, end_color=BRAND_SOFT)
        key.alignment = Alignment(horizontal="right", vertical="center", indent=1)
        key.border = CELL_BORDER
        val = ws.cell(row=row, column=2)
        set_value(val, value, "datetime" if isinstance(value, datetime) else ("int" if isinstance(value, int) else "text"))
        val.font = _font()
        val.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        val.border = CELL_BORDER
        ws.row_dimensions[row].height = 24
        row += 1
    row += 1

    for paragraph in paragraphs:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
        cell = ws.cell(row=row, column=1)
        set_text(cell, paragraph)
        cell.font = _font(color="1F2D29")
        cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
        total_width = sum(widths[:span])
        ws.row_dimensions[row].height = min(max(22, 17 * math.ceil(len(paragraph) / max(total_width * 0.95, 20)) + 6), 140)
        row += 1
    row += 1

    if table_rows:
        for index, header in enumerate(table_header, start=1):
            cell = ws.cell(row=row, column=index)
            set_text(cell, header)
            cell.font = _font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", start_color=BRAND, end_color=BRAND)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = CELL_BORDER
        ws.row_dimensions[row].height = 26
        row += 1
        for table_row in table_rows:
            tallest = 1
            for index, value in enumerate(table_row, start=1):
                cell = ws.cell(row=row, column=index)
                set_text(cell, value)
                cell.font = _font(bold=(index == 1))
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                cell.border = CELL_BORDER
                width = widths[min(index - 1, len(widths) - 1)]
                tallest = max(tallest, math.ceil(len(str(value)) / max(width * 0.9, 8)))
            ws.row_dimensions[row].height = min(max(22, 16 * tallest + 6), 160)
            row += 1


def new_workbook(title: str) -> Workbook:
    wb = Workbook()
    wb.properties.creator = "راستی‌سی"
    wb.properties.title = title
    wb.properties.keywords = "RastiSi"
    return wb


def workbook_to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ───────────────────────────── اعتبارسنجیِ آپلود ─────────────────────────────

def upload_extension(name: str) -> str:
    lowered = (name or "").lower().strip()
    dot = lowered.rfind(".")
    return lowered[dot:] if dot != -1 else ""


def validate_import_upload(uploaded_file) -> str:
    """بررسیِ امنیتیِ فایلِ آپلودیِ واردات؛ پسوندِ پذیرفته‌شده (``.xlsx`` یا
    ``.csv``) را برمی‌گرداند و در غیرِ این صورت ``XlsxUploadError`` می‌اندازد.
    به نام/نوعِ اعلام‌شده بسنده نمی‌کند: برای XLSX ساختارِ ZIP را واقعاً
    می‌خواند."""
    name = getattr(uploaded_file, "name", "") or ""
    if "/" in name or "\\" in name or "\x00" in name:
        raise XlsxUploadError("نامِ فایل نامعتبر است.")
    extension = upload_extension(name)
    if extension in REJECTED_UPLOAD_EXTENSIONS:
        raise XlsxUploadError(
            "این نوعِ فایل پذیرفته نمی‌شود. فایلِ اکسل را با قالبِ «Excel Workbook (*.xlsx)» ذخیره کنید؛ "
            "فایل‌هایِ ماکرودار (xlsm) و قالب‌هایِ قدیمی (xls) مجاز نیستند."
        )
    if extension not in (".xlsx", ".csv"):
        raise XlsxUploadError("فقط فایلِ اکسل (xlsx) پذیرفته می‌شود.")

    size = getattr(uploaded_file, "size", None)
    if size is not None and size > MAX_IMPORT_FILE_SIZE_BYTES:
        raise XlsxUploadError(
            f"حجمِ فایل نباید بیشتر از {MAX_IMPORT_FILE_SIZE_BYTES // (1024 * 1024)} مگابایت باشد."
        )
    content_type = getattr(uploaded_file, "content_type", "") or ""
    if content_type and content_type not in ALLOWED_UPLOAD_CONTENT_TYPES:
        raise XlsxUploadError(f"نوعِ فایلِ «{content_type}» پذیرفته نمی‌شود.")

    if extension == ".xlsx":
        _validate_xlsx_container(uploaded_file)
    return extension


def _validate_xlsx_container(uploaded_file) -> None:
    uploaded_file.seek(0)
    try:
        if not zipfile.is_zipfile(uploaded_file):
            raise XlsxUploadError("این فایل یک فایلِ اکسلِ معتبر (xlsx) نیست یا خراب است.")
        uploaded_file.seek(0)
        with zipfile.ZipFile(uploaded_file) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_ENTRIES:
                raise XlsxUploadError("ساختارِ فایلِ اکسل نامعتبر است.")
            if sum(info.file_size for info in infos) > MAX_UNCOMPRESSED_BYTES:
                raise XlsxUploadError("محتوایِ فایلِ اکسل بیش از حدِ مجاز بزرگ است.")
            names = {info.filename.lower() for info in infos}
            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise XlsxUploadError("این فایل یک فایلِ اکسلِ معتبر (xlsx) نیست.")
            if any(name.endswith("vbaproject.bin") or "/vbaproject" in name for name in names):
                raise XlsxUploadError("فایل‌هایِ دارایِ ماکرو پذیرفته نمی‌شوند.")
            content_types = archive.read("[Content_Types].xml").decode("utf-8", "ignore").lower()
            if "macroenabled" in content_types or "vbaproject" in content_types:
                raise XlsxUploadError("فایل‌هایِ دارایِ ماکرو پذیرفته نمی‌شوند.")
            if "spreadsheetml.sheet.main+xml" not in content_types:
                raise XlsxUploadError("این فایل یک فایلِ اکسلِ معتبر (xlsx) نیست.")
    except zipfile.BadZipFile as exc:
        raise XlsxUploadError("این فایل یک فایلِ اکسلِ معتبر (xlsx) نیست یا خراب است.") from exc
    finally:
        uploaded_file.seek(0)


# ───────────────────────────── خواندنِ XLSX ─────────────────────────────

@dataclass
class XlsxRow:
    number: int  # شماره‌ی ردیف در خودِ اکسل (۱-مبنا، هدر = ۱)
    cells: list  # مقدارهایِ متنیِ سلول‌ها به ترتیبِ ستون
    formula_columns: list = field(default_factory=list)  # اندیسِ ستون‌هایِ فرمولی/خطادار


class XlsxReadError(Exception):
    """فایل خوانده نشد (پیامِ آن برایِ نمایشِ مستقیم امن است)."""


class XlsxRowLimitExceededError(Exception):
    pass


def cell_to_text(value) -> str:
    """مقدارِ خامِ یک سلول را به رشته‌ی ورودیِ قابلِ‌پردازش تبدیل می‌کند."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)).normalize(), "f")
    if isinstance(value, Decimal):
        return format(value.normalize(), "f")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).replace("\x00", "")
    text = _ILLEGAL_XML_RE.sub("", text)
    return text[:MAX_FIELD_LENGTH]


def read_xlsx_table(fileobj, *, sheet_names: tuple = ("داده‌ها", "داده ها", "Data"),
                    max_rows: int = MAX_IMPORT_ROWS):
    """جدولِ اصلیِ یک XLSX را می‌خواند: ``(headers, rows)``. شیتِ «داده‌ها» اگر
    باشد، وگرنه نخستین شیتِ غیرِ راهنما. فرمول‌ها اجرا نمی‌شوند (``data_only``
    روشن نیست و سلولِ فرمولی فقط علامت‌گذاری می‌شود). حداکثرِ ردیف رعایت
    می‌شود و ردیف‌هایِ کاملاً خالی نادیده گرفته می‌شوند."""
    try:
        workbook = load_workbook(fileobj, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:  # noqa: BLE001 — openpyxl انواعِ مختلفِ استثنا می‌اندازد
        raise XlsxReadError("فایلِ اکسل خوانده نشد؛ ممکن است خراب باشد.") from exc
    try:
        sheet = None
        by_name = {name.strip(): name for name in workbook.sheetnames}
        for wanted in sheet_names:
            if wanted in by_name:
                sheet = workbook[by_name[wanted]]
                break
        if sheet is None:
            skip = {"راهنما", "فهرست‌ها", "فهرست ها", "درباره فایل"}
            for name in workbook.sheetnames:
                if name.strip() not in skip:
                    sheet = workbook[name]
                    break
        if sheet is None:
            raise XlsxReadError("هیچ شیتِ دادهای در فایل پیدا نشد.")
        try:
            sheet.reset_dimensions()
        except Exception:  # noqa: BLE001
            pass

        headers: list[str] = []
        rows: list[XlsxRow] = []
        blank_streak = 0
        for row_index, excel_row in enumerate(sheet.iter_rows(max_col=MAX_COLUMNS), start=1):
            if not headers:
                values = [cell_to_text(getattr(c, "value", None)).strip() for c in excel_row]
                if any(values):
                    while values and not values[-1]:
                        values.pop()
                    headers = values
                elif row_index > 20:
                    break
                continue
            texts, formulas = [], []
            for col_index, cell in enumerate(excel_row[: len(headers)]):
                data_type = getattr(cell, "data_type", "n")
                if data_type in ("f", "e"):
                    formulas.append(col_index)
                    texts.append("")
                    continue
                texts.append(cell_to_text(getattr(cell, "value", None)))
            if not any(t.strip() for t in texts) and not formulas:
                blank_streak += 1
                if blank_streak > MAX_TRAILING_BLANK_ROWS:
                    break
                continue
            blank_streak = 0
            if len(rows) >= max_rows:
                raise XlsxRowLimitExceededError(f"تعدادِ ردیف‌ها از حداکثرِ مجاز ({max_rows}) بیشتر است.")
            texts += [""] * (len(headers) - len(texts))
            rows.append(XlsxRow(number=row_index, cells=texts, formula_columns=formulas))
        return headers, rows
    finally:
        workbook.close()
