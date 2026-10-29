"""Builds the QA sample workbooks from the REAL store-specific template (same engine/template the product list button downloads).

    python manage.py shell < docs/admin-upgrade/samples/make_samples.py       (run from the repo root; QA database only)

Files (written next to this script):
  01_five_new_brands_and_products.xlsx   5 brand-new brands + 8 products, one new subcategory under an existing main category,
                                         names written with spacing/case variants of EXISTING brands (must NOT become proposals)
  02_update_existing_by_sku.xlsx         update-only rows (SKU + price + stock; every other column blank = unchanged)
  03_mixed_errors_and_new_main_category.xlsx   a new main category + subcategory, specs/attributes, and deliberately bad rows
"""
import io
from pathlib import Path

from openpyxl import load_workbook

from apps.dashboard.services import import_xlsx
from apps.stores.models import Store

OUT = Path("docs/admin-upgrade/samples")
store = Store.objects.get(slug="rasti-mode-demo")
COLS = [c for c in import_xlsx.IMPORT_SPECS["products"] if c.in_template]
KEYS = [c.key for c in COLS]


def build(name, rows):
    wb = load_workbook(io.BytesIO(import_xlsx.build_template_xlsx("products", store)))
    ws = wb["داده‌ها"]
    for r, data in enumerate(rows, start=2):
        for key, value in data.items():
            cell = ws.cell(row=r, column=KEYS.index(key) + 1)
            cell.value = value
    wb.save(OUT / name)
    print("wrote", name, len(rows), "rows")


NEW_BRANDS = ["آروین", "بامداد", "پارسا", "تندیس", "ژیوار"]
build("01_five_new_brands_and_products.xlsx", [
    {"name": "کتانی رانینگ آروین", "category_code": "کفش", "subcategory_name": "کتانی رانینگ", "price": 1850000, "sku": "QA-NB-01", "brand_code": NEW_BRANDS[0], "stock": 20, "specifications": "جنس رویه: مش؛ وزن: ۲۸۰ گرم"},
    {"name": "کتانی کژوال بامداد", "category_code": "کفش", "subcategory_name": "کتانی کژوال", "price": 1450000, "sku": "QA-NB-02", "brand_code": NEW_BRANDS[1], "stock": 12},
    {"name": "شلوار جین پارسا", "category_code": "پوشاک", "subcategory_name": "شلوار جین", "price": 990000, "sku": "QA-NB-03", "brand_code": NEW_BRANDS[2], "stock": 35, "specifications": "جنس: جین ۱۲ اونس"},
    {"name": "کاپشن تندیس", "category_code": "پوشاک", "subcategory_name": "کاپشن و بامبر", "price": 3250000, "sku": "QA-NB-04", "brand_code": NEW_BRANDS[3], "stock": 8},
    {"name": "کیف دوشی ژیوار", "category_code": "کیف", "subcategory_name": "کیف دوشی و مجلسی", "price": 2100000, "sku": "QA-NB-05", "brand_code": NEW_BRANDS[4], "stock": 15},
    {"name": "کیف پول چرمی", "category_code": "کیف", "subcategory_name": "کیف پول و جاکارتی", "price": 650000, "sku": "QA-NB-06", "brand_code": NEW_BRANDS[4], "stock": 40},
    {"name": "شلوار کژوال موشن", "category_code": "پوشاک > شلوار کژوال", "price": 870000, "sku": "QA-NB-07", "brand_code": "demo  motion", "stock": 22},
    {"name": "ژاکت اربن", "category_code": "پوشاک", "subcategory_name": "ژاکت چرم و اورشرت", "price": 2750000, "sku": "QA-NB-08", "brand_code": "DEMO URBAN", "stock": 9},
])
build("02_update_existing_by_sku.xlsx", [
    {"sku": "QA-NB-01", "price": 1790000, "stock": 26},
    {"sku": "QA-NB-03", "price": 1050000},
    {"sku": "QA-NB-06", "stock": 55},
])
build("03_mixed_errors_and_new_main_category.xlsx", [
    {"name": "هدفون بی‌سیم", "category_code": "لوازم دیجیتال", "subcategory_name": "هدفون", "price": 2400000, "sku": "QA-MX-01", "brand_code": "آروین", "stock": 10, "attributes_text": "رنگ: مشکی؛ اتصال: بلوتوث"},
    {"name": "", "category_code": "کفش", "subcategory_name": "کتانی رانینگ", "price": "abc", "sku": "QA-MX-02"},
    {"name": "دسته‌یِ اصلیِ تنها", "category_code": "پوشاک", "price": 100000, "sku": "QA-MX-03"},
])
